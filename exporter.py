"""
Moduł eksporterów danych Facebooka do formatów: Excel (xlsx), HTML, PDF, JSON
"""

import json
from pathlib import Path
from typing import Dict, Any, Optional

# openpyxl
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# reportlab
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Rejestracja czcionki wspierającej polskie znaki dla PDF
POLISH_FONT_NAME = "Helvetica"
for font_candidate in [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
]:
    if Path(font_candidate).exists():
        try:
            pdfmetrics.registerFont(TTFont("DejaVu", font_candidate))
            POLISH_FONT_NAME = "DejaVu"
            break
        except Exception:
            pass


def export_json(data: dict, file_path: Path):
    """Eksport do pliku JSON."""
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def export_excel(data: dict, file_path: Path):
    """Generuje wieloarkuszowy, ostylowany plik Excel (.xlsx)."""
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1877F2", end_color="1877F2", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    ident = data.get("identity") or data
    stats = data.get("stats") or data.get("counts") or {}
    about = data.get("about") or {}
    posts = data.get("posts") or []
    photos = data.get("photos") or []
    videos = data.get("videos") or []
    network = data.get("network") or {}

    # --- ARKUSZ 1: Profil & Bio ---
    ws_prof = wb.active
    ws_prof.title = "Profil"
    ws_prof.append(["Klucz", "Wartość"])
    ws_prof["A1"].fill = header_fill
    ws_prof["B1"].fill = header_fill
    ws_prof["A1"].font = header_font
    ws_prof["B1"].font = header_font

    prof_rows = [
        ("Imię i nazwisko", ident.get("name")),
        ("ID", ident.get("id")),
        ("Nazwa alternatywna", ident.get("alternate_name")),
        ("Płeć", ident.get("gender")),
        ("Profil zweryfikowany", "Tak" if ident.get("is_verified") else "Nie"),
        ("Link do profilu", ident.get("url")),
        ("Zdjęcie profilowe", ident.get("profile_picture_url")),
        ("Obserwujący", stats.get("followers_count") or stats.get("followers")),
        ("Obserwowani", stats.get("followees_count") or stats.get("following")),
        ("Znajomi", stats.get("friends_count") or stats.get("friends")),
        ("Bio", about.get("biography")),
        ("Ulubione cytaty", about.get("favorite_quotes")),
        ("Kategoria", about.get("category")),
        ("Aktualna lokalizacja", about.get("current_location")),
        ("Miejscowość rodzinna", about.get("hometown")),
        ("Data urodzenia", about.get("birthday")),
        ("Związek", f"{about.get('relationship', {}).get('type','')} {about.get('relationship', {}).get('desc','')}"),
        ("Języki", ", ".join(about.get("languages", [])) if isinstance(about.get("languages"), list) else about.get("languages")),
        ("Praca", "; ".join([f"{w.get('company','')} ({w.get('desc','')})" for w in about.get("works", [])])),
        ("Edukacja (Uczelnie)", "; ".join([c.get("name","") for c in about.get("colleges", [])])),
        ("Edukacja (Szkoły)", "; ".join([h.get("name","") for h in about.get("high_schools", [])])),
        ("Rodzina", "; ".join([f"{f.get('name','')}: {f.get('type','')}" for f in about.get("family", [])])),
    ]
    for row in prof_rows:
        ws_prof.append([row[0], str(row[1]) if row[1] is not None else ""])

    # --- ARKUSZ 2: Posty ---
    ws_posts = wb.create_sheet(title="Posty")
    post_headers = ["ID Posta", "Data Publikacji", "Treść", "Polubienia", "Komentarze", "Udostępnienia", "Link URL"]
    ws_posts.append(post_headers)
    for col_num in range(1, len(post_headers) + 1):
        c = ws_posts.cell(row=1, column=col_num)
        c.fill = header_fill
        c.font = header_font

    for p in posts:
        ws_posts.append([
            p.get("id"),
            p.get("creation_time_formatted") or p.get("creation_time"),
            p.get("text"),
            p.get("reaction_count") or p.get("reactions_count"),
            p.get("comments_count"),
            p.get("share_count"),
            p.get("url"),
        ])

    # --- ARKUSZ 3: Komentarze ---
    ws_comments = wb.create_sheet(title="Komentarze")
    com_headers = ["ID Posta", "ID Komentarza", "Autor", "Treść Komentarza", "Data", "Reakcje", "Odpowiedzi (Liczba)"]
    ws_comments.append(com_headers)
    for col_num in range(1, len(com_headers) + 1):
        c = ws_comments.cell(row=1, column=col_num)
        c.fill = header_fill
        c.font = header_font

    for p in posts:
        pid = p.get("id")
        for com in p.get("comments", []):
            author = com.get("owner", {}).get("name") if isinstance(com.get("owner"), dict) else "Anonim"
            ws_comments.append([
                pid,
                com.get("id"),
                author,
                com.get("text"),
                com.get("created_time_formatted") or com.get("created_time"),
                com.get("reaction_count"),
                len(com.get("replies", [])),
            ])
            for rep in com.get("replies", []):
                rep_author = rep.get("owner", {}).get("name") if isinstance(rep.get("owner"), dict) else "Anonim"
                ws_comments.append([
                    pid,
                    rep.get("id"),
                    f"↳ Odp: {rep_author}",
                    rep.get("text"),
                    rep.get("created_time_formatted") or rep.get("created_time"),
                    rep.get("reaction_count"),
                    0,
                ])

    # --- ARKUSZ 4: Multimedia ---
    ws_media = wb.create_sheet(title="Multimedia")
    media_headers = ["Typ", "ID", "Data / Długość", "Link CDN / URL", "Wyświetlenia"]
    ws_media.append(media_headers)
    for col_num in range(1, len(media_headers) + 1):
        c = ws_media.cell(row=1, column=col_num)
        c.fill = header_fill
        c.font = header_font

    for ph in photos:
        ws_media.append(["Zdjęcie", ph.get("id"), "", ph.get("image_url") or ph.get("url"), ""])

    for v in videos:
        length_str = f"{v.get('length_in_second')}s" if v.get("length_in_second") else ""
        ws_media.append(["Wideo/Rolka", v.get("id"), length_str, v.get("video_url") or v.get("url"), v.get("play_count")])

    # --- ARKUSZ 5: Sieć kontaktów ---
    ws_net = wb.create_sheet(title="Sieć kontaktów")
    net_headers = ["Relacja", "ID", "Imię i nazwisko", "Profil URL", "Zweryfikowany"]
    ws_net.append(net_headers)
    for col_num in range(1, len(net_headers) + 1):
        c = ws_net.cell(row=1, column=col_num)
        c.fill = header_fill
        c.font = header_font

    for u in network.get("friends", []):
        if u:
            ws_net.append(["Znajomy", u.get("id"), u.get("name"), u.get("url"), "Tak" if u.get("is_verified") else "Nie"])

    for u in network.get("followers", []):
        if u:
            ws_net.append(["Obserwujący", u.get("id"), u.get("name"), u.get("url"), "Tak" if u.get("is_verified") else "Nie"])

    for u in network.get("followees", []):
        if u:
            ws_net.append(["Obserwowany", u.get("id"), u.get("name"), u.get("url"), "Tak" if u.get("is_verified") else "Nie"])

    # Automatyczne dopasowanie szerokości kolumn
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    wb.save(file_path)


def export_html(data: dict, file_path: Path):
    """Generuje interaktywny, nowoczesny dashboard HTML."""
    ident = data.get("identity") or data
    stats = data.get("stats") or data.get("counts") or {}
    about = data.get("about") or {}
    posts = data.get("posts") or []
    photos = data.get("photos") or []
    videos = data.get("videos") or []
    network = data.get("network") or {}

    avatar_url = ident.get("profile_picture_url") or "https://via.placeholder.com/150"
    name = ident.get("name") or "Nieznany profil"
    uid = ident.get("id") or ""
    verified_badge = '<span class="badge verified">✓ Zweryfikowany</span>' if ident.get("is_verified") else ""

    # Generowanie HTML postów
    posts_html = ""
    for idx, p in enumerate(posts, start=1):
        comments_html = ""
        for c in p.get("comments", []):
            author = c.get("owner", {}).get("name") if isinstance(c.get("owner"), dict) else "Anonim"
            comments_html += f"""
            <div class="comment">
                <strong>{author}</strong> <span class="comment-date">({c.get('created_time_formatted') or ''})</span>:
                <p>{c.get('text') or ''}</p>
            </div>
            """
            for rep in c.get("replies", []):
                rep_author = rep.get("owner", {}).get("name") if isinstance(rep.get("owner"), dict) else "Anonim"
                comments_html += f"""
                <div class="comment reply">
                    ↳ <strong>{rep_author}</strong>: {rep.get('text') or ''}
                </div>
                """

        attachments_html = ""
        for att in p.get("attachments", []):
            img = att.get("image_url")
            if img:
                attachments_html += f'<img src="{img}" class="post-img" loading="lazy" alt="Załącznik">'

        posts_html += f"""
        <div class="post-card">
            <div class="post-header">
                <div>
                    <span class="post-number">Post #{idx}</span>
                    <span class="post-date">{p.get('creation_time_formatted') or ''}</span>
                </div>
                <a href="{p.get('url') or '#'}" target="_blank" class="post-link">Otwórz na FB ↗</a>
            </div>
            <div class="post-body">
                <p>{(p.get('text') or '').replace(chr(10), '<br>')}</p>
                {attachments_html}
            </div>
            <div class="post-stats">
                <span>👍 {p.get('reaction_count') or p.get('reactions_count') or 0}</span>
                <span>💬 {p.get('comments_count') or 0}</span>
                <span>🔄 {p.get('share_count') or 0}</span>
            </div>
            {f'<div class="comments-section"><h4>Komentarze ({len(p.get("comments", []))})</h4>{comments_html}</div>' if comments_html else ''}
        </div>
        """

    # HTML mediów
    media_html = ""
    for ph in photos:
        img_url = ph.get("image_url") or ph.get("url")
        ph_url = ph.get("url") or "#"
        if img_url:
            media_html += f'<div class="media-thumb"><a href="{ph_url}" target="_blank" title="Otwórz zdjęcie na FB"><img src="{img_url}" loading="lazy"></a></div>'

    # HTML znajomych
    friends = network.get("friends") or []
    friends_html = ""
    for fr in friends:
        fr_avatar = fr.get("profile_picture_url") or "https://via.placeholder.com/60"
        fr_name = fr.get("name") or "Użytkownik"
        fr_url = fr.get("url") or "#"
        fr_id = fr.get("id") or ""
        friends_html += f"""
        <div class="friend-card">
            <img src="{fr_avatar}" class="friend-avatar" loading="lazy" alt="Avatar">
            <div class="friend-info">
                <a href="{fr_url}" target="_blank" class="friend-name">{fr_name}</a>
                <span class="friend-id">ID: {fr_id}</span>
            </div>
        </div>
        """

    html_content = f"""<!DOCTYPE html>
<html lang="pl">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Raport OSINT: {name}</title>
    <style>
        :root {{
            --bg: #0f172a;
            --surface: #1e293b;
            --surface-hover: #334155;
            --primary: #3b82f6;
            --accent: #10b981;
            --text: #f8fafc;
            --text-muted: #94a3b8;
            --border: #334155;
            --radius: 12px;
        }}
        * {{ box-sizing: border-box; margin: 0; padding: 0; }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            background: var(--bg);
            color: var(--text);
            padding: 24px;
            line-height: 1.5;
        }}
        .container {{ max-width: 1100px; margin: 0 auto; }}
        .header-card {{
            background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%);
            border: 1px solid var(--border);
            border-radius: var(--radius);
            padding: 24px;
            display: flex;
            gap: 24px;
            align-items: center;
            margin-bottom: 24px;
        }}
        .avatar {{ width: 110px; height: 110px; border-radius: 50%; object-fit: cover; border: 3px solid var(--primary); }}
        .header-info h1 {{ font-size: 26px; font-weight: 700; margin-bottom: 6px; }}
        .header-info p {{ color: var(--text-muted); font-size: 14px; margin-bottom: 4px; }}
        .badge {{ display: inline-block; padding: 3px 10px; border-radius: 999px; font-size: 12px; font-weight: 600; margin-top: 6px; }}
        .badge.verified {{ background: rgba(59, 130, 246, 0.2); color: var(--primary); border: 1px solid var(--primary); }}
        .stats-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); gap: 16px; margin-bottom: 24px; }}
        .stat-card {{ background: var(--surface); border: 1px solid var(--border); border-radius: var(--radius); padding: 18px; text-align: center; }}
        .stat-val {{ font-size: 24px; font-weight: 700; color: var(--primary); }}
        .stat-label {{ color: var(--text-muted); font-size: 13px; text-transform: uppercase; margin-top: 4px; }}
        .section-title {{ font-size: 20px; font-weight: 700; margin: 28px 0 16px 0; border-bottom: 2px solid var(--primary); padding-bottom: 6px; display: inline-block; }}
        .about-card {{ background: var(--surface); border: 1px solid var(--border); border-radius: var(--radius); padding: 20px; margin-bottom: 24px; }}
        .about-card p {{ margin-bottom: 8px; font-size: 15px; }}
        .post-card {{ background: var(--surface); border: 1px solid var(--border); border-radius: var(--radius); padding: 20px; margin-bottom: 20px; }}
        .post-header {{ display: flex; justify-content: space-between; margin-bottom: 12px; color: var(--text-muted); font-size: 13px; }}
        .post-number {{ font-weight: 700; color: var(--primary); margin-right: 8px; }}
        .post-link {{ color: var(--primary); text-decoration: none; font-size: 13px; }}
        .post-link:hover {{ text-decoration: underline; }}
        .post-body {{ font-size: 15px; margin-bottom: 16px; word-break: break-word; }}
        .post-img {{ max-width: 100%; border-radius: 8px; margin-top: 12px; max-height: 400px; object-fit: cover; }}
        .post-stats {{ display: flex; gap: 20px; font-size: 14px; font-weight: 600; color: var(--text-muted); border-top: 1px solid var(--border); padding-top: 12px; }}
        .comments-section {{ background: rgba(0,0,0,0.2); border-radius: 8px; padding: 14px; margin-top: 14px; }}
        .comments-section h4 {{ font-size: 13px; color: var(--text-muted); margin-bottom: 10px; text-transform: uppercase; }}
        .comment {{ font-size: 13px; margin-bottom: 8px; border-left: 2px solid var(--primary); padding-left: 8px; }}
        .comment.reply {{ margin-left: 16px; border-left-color: var(--accent); }}
        .comment-date {{ color: var(--text-muted); font-size: 11px; }}
        .media-grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(130px, 1fr)); gap: 10px; margin-bottom: 24px; }}
        .media-thumb img {{ width: 100%; height: 130px; object-fit: cover; border-radius: 8px; border: 1px solid var(--border); transition: transform 0.2s; }}
        .media-thumb img:hover {{ transform: scale(1.03); border-color: var(--primary); }}
        .friends-grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(240px, 1fr)); gap: 12px; margin-bottom: 24px; }}
        .friend-card {{ background: var(--surface); border: 1px solid var(--border); border-radius: var(--radius); padding: 10px 14px; display: flex; align-items: center; gap: 12px; }}
        .friend-avatar {{ width: 44px; height: 44px; border-radius: 50%; object-fit: cover; border: 1px solid var(--border); flex-shrink: 0; }}
        .friend-info {{ overflow: hidden; }}
        .friend-name {{ font-weight: 600; color: var(--text); text-decoration: none; font-size: 14px; display: block; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
        .friend-name:hover {{ color: var(--primary); }}
        .friend-id {{ font-size: 11px; color: var(--text-muted); display: block; }}
    </style>
</head>
<body>
    <div class="container">
        <div class="header-card">
            <img src="{avatar_url}" alt="Avatar" class="avatar">
            <div class="header-info">
                <h1>{name} {verified_badge}</h1>
                <p><strong>ID:</strong> {uid} | <strong>Płeć:</strong> {ident.get('gender') or 'Nieokreślona'}</p>
                <p><strong>URL:</strong> <a href="{ident.get('url') or '#'}" target="_blank" style="color:var(--primary);">{ident.get('url')}</a></p>
                <p><em>Zrzut wygenerowany: {data.get('scraped_at')}</em></p>
            </div>
        </div>

        <div class="stats-grid">
            <div class="stat-card">
                <div class="stat-val">{stats.get('friends_count') or len(friends) or '—'}</div>
                <div class="stat-label">Znajomi</div>
            </div>
            <div class="stat-card">
                <div class="stat-val">{stats.get('followers_count') or len(network.get('followers', [])) or '—'}</div>
                <div class="stat-label">Obserwujący</div>
            </div>
            <div class="stat-card">
                <div class="stat-val">{len(posts)}</div>
                <div class="stat-label">Pobranych Postów</div>
            </div>
            <div class="stat-card">
                <div class="stat-val">{len(photos) + len(videos)}</div>
                <div class="stat-label">Multimediów</div>
            </div>
        </div>

        <h2 class="section-title">O Profilu</h2>
        <div class="about-card">
            {f'<p><strong>Biogram:</strong> {about.get("biography")}</p>' if about.get("biography") else ''}
            {f'<p><strong>Ulubione cytaty:</strong> <em>{about.get("favorite_quotes")}</em></p>' if about.get("favorite_quotes") else ''}
            {f'<p><strong>Lokalizacja:</strong> {about.get("current_location")}</p>' if about.get("current_location") else ''}
            {f'<p><strong>Miejscowość rodzinna:</strong> {about.get("hometown")}</p>' if about.get("hometown") else ''}
            {f'<p><strong>Data urodzenia:</strong> {about.get("birthday")}</p>' if about.get("birthday") else ''}
            {f'<p><strong>Związek:</strong> {about.get("relationship", {}).get("type","")} {about.get("relationship", {}).get("desc","")}</p>' if about.get("relationship", {}).get("type") else ''}
            {f'<p><strong>Praca:</strong> {", ".join([w.get("company","") for w in about.get("works", [])])}</p>' if about.get("works") else ''}
            {f'<p><strong>Edukacja (Uczelnie):</strong> {"; ".join([c.get("name","") for c in about.get("colleges", [])])}</p>' if about.get("colleges") else ''}
            {f'<p><strong>Edukacja (Szkoły średnie):</strong> {"; ".join([h.get("name","") for h in about.get("high_schools", [])])}</p>' if about.get("high_schools") else ''}
            {f'<p><strong>Rodzina:</strong> {"; ".join([f"{f.get("name","")}: {f.get("type","")}" for f in about.get("family", [])])}</p>' if about.get("family") else ''}
            {f'<p><strong>Języki:</strong> {", ".join(about.get("languages", [])) if isinstance(about.get("languages"), list) else about.get("languages")}</p>' if about.get("languages") else ''}
        </div>

        {f'<h2 class="section-title">Galeria Zdjęć ({len(photos)})</h2><div class="media-grid">{media_html}</div>' if media_html else ''}

        {f'<h2 class="section-title">Znajomi ({len(friends)})</h2><div class="friends-grid">{friends_html}</div>' if friends_html else ''}

        <h2 class="section-title">Pobrane Posty ({len(posts)})</h2>
        {posts_html if posts_html else '<p style="color:var(--text-muted);">Brak dostępnych postów.</p>'}
    </div>
</body>
</html>"""

    with open(file_path, "w", encoding="utf-8") as f:
        f.write(html_content)


def export_pdf(data: dict, file_path: Path):
    """Generuje dokument PDF z użyciem ReportLab."""
    doc = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Heading1"],
        fontName=POLISH_FONT_NAME,
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1877F2"),
        spaceAfter=10,
    )

    h2_style = ParagraphStyle(
        "H2Style",
        parent=styles["Heading2"],
        fontName=POLISH_FONT_NAME,
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=14,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontName=POLISH_FONT_NAME,
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
    )

    bold_style = ParagraphStyle(
        "BoldStyle",
        parent=body_style,
        fontName=POLISH_FONT_NAME,
        fontSize=9,
        leading=13,
        fontStyle="Bold",
    )

    ident = data.get("identity") or data
    stats = data.get("stats") or data.get("counts") or {}
    about = data.get("about") or {}
    posts = data.get("posts") or []

    story = []

    # Nagłówek
    story.append(Paragraph(f"Raport Profilu Facebook: {ident.get('name')}", title_style))
    story.append(Paragraph(f"Data zrzutu: {data.get('scraped_at')} | ID: {ident.get('id')}", body_style))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1877F2"), spaceAfter=12))

    # Tabela tożsamości
    table_data = [
        [Paragraph("Pole", bold_style), Paragraph("Wartość", bold_style)],
        [Paragraph("Imię i nazwisko", body_style), Paragraph(str(ident.get("name")), body_style)],
        [Paragraph("ID profilu", body_style), Paragraph(str(ident.get("id")), body_style)],
        [Paragraph("Płeć", body_style), Paragraph(str(ident.get("gender") or "Brak"), body_style)],
        [Paragraph("Status", body_style), Paragraph("Zweryfikowany" if ident.get("is_verified") else "Zwykły", body_style)],
        [Paragraph("Obserwujący", body_style), Paragraph(str(stats.get("followers_count") or stats.get("followers") or "Ukryte"), body_style)],
        [Paragraph("Znajomi", body_style), Paragraph(str(stats.get("friends_count") or stats.get("friends") or "Ukryte"), body_style)],
        [Paragraph("Lokalizacja", body_style), Paragraph(str(about.get("current_location") or "Brak"), body_style)],
        [Paragraph("Pochodzenie", body_style), Paragraph(str(about.get("hometown") or "Brak"), body_style)],
    ]

    t = Table(table_data, colWidths=[140, 380])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
    ]))
    story.append(t)
    story.append(Spacer(1, 14))

    # Posty
    story.append(Paragraph(f"Pobrane Posty ({len(posts)})", h2_style))

    for idx, p in enumerate(posts[:20], start=1):
        txt = (p.get("text") or "Brak tekstu").replace("\n", " ")[:200]
        p_date = p.get("creation_time_formatted") or ""
        p_stats = f"Reakcje: {p.get('reaction_count') or 0} | Komentarze: {p.get('comments_count') or 0} | Udostępnienia: {p.get('share_count') or 0}"

        post_box = [
            [Paragraph(f"Post #{idx} — {p_date} (ID: {p.get('id')})", bold_style)],
            [Paragraph(f"<i>{txt}...</i>", body_style)],
            [Paragraph(p_stats, body_style)],
        ]
        pt = Table(post_box, colWidths=[520])
        pt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f8fafc")),
            ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(pt)
        story.append(Spacer(1, 6))

    doc.build(story)
